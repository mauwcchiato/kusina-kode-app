"""Builds KusinaKode_Timetable.xlsx in the adviser's template format.

No shebang on purpose: Windows' `py` launcher honours it, and
`#!/usr/bin/env python` resolves to the Microsoft Store stub rather than the
real interpreter, so `py build_timetable.py` fails while `py -c` works.

Three tabs, matching IAN-RG_TIMETABLE_CHECKING so the format is familiar:

  Progress Table        every feature, its status, date and remarks
  Progress              a dashboard of per-module completion, by formula
  Integration and Unit  unit tests (left block) and integration tests (right)

Sources, none of them typed by hand:
  features.csv                        the feature list and status (edit this)
  api-tests.csv                       integration test results
  app/build/test-results/.../*.xml    unit results, parsed from Gradle

Deliberately a NEW workbook rather than an edit of the adviser's file: that
file holds another group's tracker plus sample templates, and its Progress
formulas already carry #REF! errors. Writing into it would destroy their work
and inherit their breakage.

Usage:
    py -m pip install openpyxl        # once
    ..\\..\\gradlew :app:testDebugUnitTest
    py build_timetable.py
"""
import csv
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
RESULTS = os.environ.get(
    "KK_TEST_RESULTS",
    os.path.join(REPO, "app", "build", "test-results", "testDebugUnitTest"),
)
OUT = os.path.join(HERE, "KusinaKode_Timetable.xlsx")

RESEARCH_QUESTION = (
    "How can a blockchain-based, gamified Filipino culinary word-puzzle "
    "application support vocabulary and heritage learning while maintaining "
    "engagement through gamification and securing rewards through a "
    "transparent, tamper-resistant blockchain mechanism?"
)

OBJECTIVES = {
    "Objective 1": "Objective 1: Provide a module that will support culinary vocabulary and "
                   "recipe-based learning and cultural awareness through interactive word puzzles.",
    "Objective 2": "Objective 2: Design a platform that will encourage user motivation, engagement "
                   "and continued participation by integrating gamification principles.",
    "Objective 3": "Objective 3: Apply blockchain technology to securely manage user rewards, "
                   "achievements and activity records in a transparent and verifiable manner.",
    "Supports 2 and 3": "Supports Objectives 2 and 3 (not tied to a single objective)",
    "Supports all objectives": "Supports all objectives (platform-wide)",
}

# ---------------------------------------------------------------- house style
ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="FFCC6B1F")
OK_FILL = PatternFill("solid", fgColor="FFE3F2E1")
WIP_FILL = PatternFill("solid", fgColor="FFFDF0D5")
NO_FILL = PatternFill("solid", fgColor="FFF2F2F2")
FLAG_FILL = PatternFill("solid", fgColor="FFFBE0E0")
MOD_FILL = PatternFill("solid", fgColor="FFF1E6D2")
THIN = Side(style="thin", color="FFBFB2A4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    "OK": OK_FILL,
    "In Progress": WIP_FILL,
    "Not Started": NO_FILL,
    "Scope Change": FLAG_FILL,
}


def body(cell, bold=False, wrap=True, size=11):
    cell.font = Font(name=ARIAL, size=size, bold=bold)
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    cell.border = BOX
    return cell


def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        body(c, bold=True)
        c.font = Font(name=ARIAL, size=11, bold=True, color="FFFFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


# ------------------------------------------------------------------ the data
def read_features():
    path = os.path.join(HERE, "features.csv")
    if not os.path.exists(path):
        sys.exit("features.csv is missing - it belongs next to this script.")
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def read_api_tests():
    with open(os.path.join(HERE, "api-tests.csv"), encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def read_unit_tests():
    if not os.path.isdir(RESULTS):
        sys.exit(
            "No unit results at %s\nRun the suite first:  gradlew :app:testDebugUnitTest" % RESULTS
        )
    rows = []
    for name in sorted(os.listdir(RESULTS)):
        if not name.endswith(".xml"):
            continue
        xml = open(os.path.join(RESULTS, name), encoding="utf-8").read()
        m = re.search(r'<testsuite name="([^"]+)"', xml)
        suite = m.group(1).split(".")[-1] if m else name
        for tc in re.finditer(r'<testcase name="([^"]+)"[^>]*time="[^"]*"\s*(/)?>', xml):
            verdict = "Pass"
            if not tc.group(2):
                tail = xml[tc.start(): xml.find("</testcase>", tc.start())]
                if "<failure" in tail:
                    verdict = "Fail"
                elif "<skipped" in tail:
                    verdict = "Skipped"
            rows.append((suite, tc.group(1).rstrip('"'), verdict))
    rows.sort(key=lambda r: (r[0], r[1]))
    return rows


# ------------------------------------------------------------------ tab 1/3
def sheet_progress_table(wb, features):
    ws = wb.create_sheet("Progress Table")
    header_row(
        ws,
        1,
        ["Problem Statement", "Objectives", "Scopes", "Features", "Platform", "Status",
         "Target Date of Completion", "Date of Completion", "Owner", "Remarks"],
        [34, 40, 26, 46, 17, 13, 17, 17, 14, 60],
    )

    row = 2
    first_data_row = row
    module_spans = {}     # module -> (first_row, last_row)
    objective_spans = {}
    for feat in features:
        obj = feat["Objective"]
        mod = feat["Module"]
        status = feat["Status"].strip()

        body(ws.cell(row=row, column=2, value=OBJECTIVES.get(obj, obj)))
        body(ws.cell(row=row, column=3, value=mod))
        body(ws.cell(row=row, column=4, value=feat["Feature"]))

        body(ws.cell(row=row, column=5, value=feat.get("Platform") or ""))

        sc = body(ws.cell(row=row, column=6, value=status))
        sc.fill = STATUS_FILL.get(status, NO_FILL)
        sc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        body(ws.cell(row=row, column=7, value=feat.get("Target Date") or ""))
        body(ws.cell(row=row, column=8, value=feat.get("Date Completed") or ""))
        body(ws.cell(row=row, column=9, value=feat.get("Owner") or ""))
        body(ws.cell(row=row, column=10, value=feat.get("Remarks") or ""))

        lo, hi = module_spans.get(mod, (row, row))
        module_spans[mod] = (min(lo, row), max(hi, row))
        lo, hi = objective_spans.get(obj, (row, row))
        objective_spans[obj] = (min(lo, row), max(hi, row))
        row += 1

    last = row - 1
    # The research question spans the whole table, as in the template.
    ws.merge_cells(start_row=first_data_row, start_column=1, end_row=last, end_column=1)
    body(ws.cell(row=first_data_row, column=1, value=RESEARCH_QUESTION))
    ws.cell(row=first_data_row, column=1).alignment = Alignment(wrap_text=True, vertical="center")

    for obj, (lo, hi) in objective_spans.items():
        if hi > lo:
            ws.merge_cells(start_row=lo, start_column=2, end_row=hi, end_column=2)
            ws.cell(row=lo, column=2).alignment = Alignment(wrap_text=True, vertical="center")
    for mod, (lo, hi) in module_spans.items():
        if hi > lo:
            ws.merge_cells(start_row=lo, start_column=3, end_row=hi, end_column=3)
        c = ws.cell(row=lo, column=3)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.font = Font(name=ARIAL, size=11, bold=True)
        c.fill = MOD_FILL

    ws.freeze_panes = "E2"
    return ws, module_spans, first_data_row, last


# ------------------------------------------------------------------ tab 2/3
def sheet_progress(wb, module_spans, table_last):
    """Dashboard. Every number is a formula over 'Progress Table', not a
    Python-computed total, so editing a status there updates this on open."""
    ws = wb.create_sheet("Progress")

    t = body(ws.cell(row=1, column=1, value="Kusina Kode - Feature Completion Tracker"), bold=True, size=14)
    t.border = Border()
    s = body(ws.cell(row=2, column=1, value="Per-module readiness. Figures are formulas over the Progress Table tab."))
    s.border = Border()
    s.font = Font(name=ARIAL, size=10, italic=True)

    header_row(ws, 4, ["Total Features", "Done", "In Progress", "Not Started", "Needs Decision", "Overall %"],
               [17, 13, 14, 14, 16, 12])
    q = "'Progress Table'!$F$2:$F$%d" % table_last
    for col, formula in enumerate([
        '=COUNTA(%s)' % q,
        '=COUNTIF(%s,"OK")' % q,
        '=COUNTIF(%s,"In Progress")' % q,
        '=COUNTIF(%s,"Not Started")' % q,
        '=COUNTIF(%s,"Scope Change")' % q,
        '=IFERROR(B6/A6,0)',
    ], start=1):
        c = body(ws.cell(row=6, column=col, value=formula), bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if col == 6:
            c.number_format = "0.0%"

    header_row(ws, 8, ["#", "Module", "Total Features", "Done", "% Complete", "Status"],
               [6, 40, 15, 10, 13, 16])

    row = 9
    ordered = sorted(module_spans.items(), key=lambda kv: kv[1][0])
    for i, (mod, (lo, hi)) in enumerate(ordered, start=1):
        rng = "'Progress Table'!$F$%d:$F$%d" % (lo, hi)
        body(ws.cell(row=row, column=1, value=i)).alignment = Alignment(horizontal="center")
        body(ws.cell(row=row, column=2, value=mod))
        body(ws.cell(row=row, column=3, value="=COUNTA(%s)" % rng)).alignment = Alignment(horizontal="center")
        body(ws.cell(row=row, column=4, value='=COUNTIF(%s,"OK")' % rng)).alignment = Alignment(horizontal="center")
        pc = body(ws.cell(row=row, column=5, value="=IFERROR(D%d/C%d,0)" % (row, row)))
        pc.number_format = "0.0%"
        pc.alignment = Alignment(horizontal="center")
        # Nested IF rather than IFS: IFS needs an _xlfn. prefix to survive.
        body(ws.cell(row=row, column=6,
                     value='=IF(E{r}=1,"Complete",IF(E{r}>0,"In Progress","Not Started"))'.format(r=row)))
        row += 1

    body(ws.cell(row=row, column=2, value="TOTAL"), bold=True)
    for col, f in ((3, "=SUM(C9:C%d)" % (row - 1)), (4, "=SUM(D9:D%d)" % (row - 1))):
        c = body(ws.cell(row=row, column=col, value=f), bold=True)
        c.alignment = Alignment(horizontal="center")
    tot = body(ws.cell(row=row, column=5, value="=IFERROR(D%d/C%d,0)" % (row, row)), bold=True)
    tot.number_format = "0.0%"
    tot.alignment = Alignment(horizontal="center")

    note = ws.cell(row=row + 2, column=1,
                   value="Status legend: OK = built and verified · In Progress = started, "
                         "incomplete · Not Started = no work yet · Scope Change = built "
                         "differently from the written feature, needs a team decision.")
    note.font = Font(name=ARIAL, size=9, italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 3, end_column=6)
    return ws


# ------------------------------------------------------------------ tab 3/3
def sheet_tests(wb, units, apis):
    ws = wb.create_sheet("Integration and Unit")
    header_row(
        ws, 1,
        ["Test No.", "Test Case", "Remarks (Pass/Fail)", "",
         "Test No.", "Test Case", "Description of Test Case", "Remarks (Pass/Fail)"],
        [11, 52, 15, 3, 11, 34, 60, 15],
    )
    ws.cell(row=1, column=4).fill = PatternFill()
    ws.cell(row=1, column=4).border = Border()

    for i, (suite, name, verdict) in enumerate(units, start=2):
        body(ws.cell(row=i, column=1, value="U-%03d" % (i - 1))).alignment = Alignment(horizontal="center")
        body(ws.cell(row=i, column=2, value="%s: %s" % (suite, name)))
        c = body(ws.cell(row=i, column=3, value=verdict))
        c.alignment = Alignment(horizontal="center")
        c.fill = OK_FILL if verdict == "Pass" else FLAG_FILL

    for i, r in enumerate(apis, start=2):
        body(ws.cell(row=i, column=5, value="I-%03d" % (i - 1))).alignment = Alignment(horizontal="center")
        body(ws.cell(row=i, column=6, value=r["Scenario"]))
        body(ws.cell(row=i, column=7,
                     value="%s | input: %s | expected: %s | actual: %s"
                           % (r["Endpoint"], r["Input"], r["Expected"], r["Actual"])))
        verdict = "Pass" if r["Result"].strip().upper() == "PASS" else "Fail"
        c = body(ws.cell(row=i, column=8, value=verdict))
        c.alignment = Alignment(horizontal="center")
        c.fill = OK_FILL if verdict == "Pass" else FLAG_FILL

    ws.freeze_panes = "A2"
    return ws


def write_status(features, units, apis):
    """Writes docs/STATUS.md — the 'where are we' brief any human or AI reads
    first. Generated, so it cannot drift from features.csv the way a
    hand-maintained summary would."""
    import datetime

    by_status = {}
    for f in features:
        by_status.setdefault(f["Status"].strip(), []).append(f)

    modules = []
    seen = []
    for f in features:
        if f["Module"] not in seen:
            seen.append(f["Module"])
    for mod in seen:
        rows = [f for f in features if f["Module"] == mod]
        ok = sum(1 for r in rows if r["Status"].strip() == "OK")
        modules.append((mod, ok, len(rows), rows[0]["Owner"] or "—"))

    defects_path = os.path.join(HERE, "defects.csv")
    open_defects = []
    if os.path.exists(defects_path):
        with open(defects_path, encoding="utf-8-sig", newline="") as fh:
            open_defects = [d for d in csv.DictReader(fh)
                            if d.get("Status", "").strip().upper() == "OPEN"]

    ok = len(by_status.get("OK", []))
    total = len(features)
    failed = sum(1 for u in units if u[2] != "Pass")

    out = os.path.join(REPO, "docs", "STATUS.md")
    with open(out, "w", encoding="utf-8", newline="\n") as fh:
        w = fh.write
        w("# Where the project stands\n\n")
        w("> Generated by `docs/testing/build_timetable.py` on %s. "
          "Do not edit by hand — change `docs/testing/features.csv` and rebuild.\n\n"
          % datetime.date.today().isoformat())

        w("**%d of %d features complete (%.0f%%)** · %d unit tests, %d failing · "
          "%d integration checks · %d open defect%s\n\n"
          % (ok, total, 100.0 * ok / total if total else 0, len(units), failed,
             len(apis), len(open_defects), "" if len(open_defects) == 1 else "s"))

        w("## By module\n\n| Module | Done | Owner |\n|---|---|---|\n")
        for mod, done, tot, owner in modules:
            mark = "done" if done == tot else ("started" if done else "not started")
            w("| %s | %d/%d — %s | %s |\n" % (mod, done, tot, mark, owner))

        for label, heading in (("In Progress", "In progress"),
                               ("Scope Change", "Needs a team decision"),
                               ("Not Started", "Not started")):
            rows = by_status.get(label, [])
            if not rows:
                continue
            w("\n## %s (%d)\n\n" % (heading, len(rows)))
            for r in rows:
                note = (r["Remarks"] or "").strip()
                w("- **%s** — %s%s\n" % (r["Feature"], r["Owner"] or "unassigned",
                                         (" · " + note) if note else ""))

        if open_defects:
            w("\n## Open defects\n\n")
            for d in open_defects:
                w("- **%s** (%s, %s) — %s\n"
                  % (d["ID"], d["Severity"], d["Module"], d["Defect"]))

        w("\n## Next actions\n\n")
        w("Derived from what is unblocked, not from a plan:\n\n")
        for r in by_status.get("Scope Change", []):
            w("1. **Decide:** %s — %s\n" % (r["Feature"], r["Remarks"]))
        for r in features:
            if r["Owner"].strip() == "UNASSIGNED":
                w("1. **Assign an owner:** %s (built, but in no signed module)\n" % r["Feature"])
        w("1. See `docs/testing/gaps.csv` for what has no test coverage.\n")

    return out


def main():
    features = read_features()
    units = read_unit_tests()
    apis = read_api_tests()

    wb = Workbook()
    wb.remove(wb.active)
    _, module_spans, _, table_last = sheet_progress_table(wb, features)
    sheet_progress(wb, module_spans, table_last)
    sheet_tests(wb, units, apis)

    # No LibreOffice here to compute the formulas, so ask Excel to do it on
    # open. Without this the dashboard cells would read blank.
    wb.calculation.fullCalcOnLoad = True

    # STATUS.md first: it is a plain file, so it lands even if the workbook is
    # open in Excel and locked.
    status_path = write_status(features, units, apis)

    try:
        wb.save(OUT)
    except PermissionError:
        print("Wrote %s" % status_path)
        sys.exit(
            "\nCould not write %s - it is open in Excel.\n"
            "Close the workbook and run this again. STATUS.md was still updated."
            % os.path.basename(OUT)
        )

    done = sum(1 for f in features if f["Status"].strip() == "OK")
    failed = sum(1 for u in units if u[2] != "Pass")
    print("Wrote %s" % OUT)
    print("Wrote %s" % status_path)
    print("  features     %d (%d OK, %d outstanding)" % (len(features), done, len(features) - done))
    print("  unit tests   %d (%d not passing)" % (len(units), failed))
    print("  integration  %d" % len(apis))
    print("  modules      %d" % len(module_spans))


if __name__ == "__main__":
    main()
